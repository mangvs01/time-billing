"""Export of the time entries as a .xlsx with the structure of the base workbook « Time sheet CP » (Decision Center billing-gsheet-export-mechanism
= DOWNLOADABLE_EXCEL, 2026-10-07): the user imports the file into Google Sheets by hand; nothing is ever written to the source Google Sheet.

One sheet per month, named like the workbook (2025: JAN, FEV, MAR, AVR, MAY, JUN, JULY, AUG, SEPT, OCT, NOV, DEC; from 2026: « JAN 26 » … « SEPT 26 »):
  B1 « Feuille de temps », B2 first day of the month, B3/B4 « Total des heures de travail du mois » = SUM of the column « Heures de travail »;
  table « Feuille de temps » from B5: Date(s) · Motif · Projet · Numéro de projet · Heure début · Heure fin · Heures de travail · Colonne1, a totals row;
  hours in DECIMAL hours (A-03: 14:09 → 14.15), « Heures de travail » = Heure fin − Heure début as in the workbook, except where that formula would not give
  the hours stored in Billing (declared duration of an imported line, a duration without times, a period across midnight): the stored hours are written;
  recap table from J6: PROJET · TOTAL H (value) · TOTAL J (= TOTAL H / hours per day) · Num projet · PO · Jours PO · J mois précédant(s) (link to the previous
  month's sheet when the mandate is there, otherwise its days before the month) · NB J TOTAL · J restants sur PO, with the workbook's formulas.
Every entry goes to the sheet of its month of service; Colonne1 stays empty (the A-01 notes are in the Motif).

SUIVI sheets (Decision Center billing-gsheet-export-scope = B, 2026-10-07): one « SUIVI 2025 », « SUIVI 2026 »… per year, after the last month sheet of
that year as in the workbook, with the columns of SUIVI 2026 for every year (amounts HT in CHF, 2025 included): PROJET · Dates · Client · Num projet
· PO · NB J PO · NB J · TARIF · TOTAL · one column per document date · Statut. The documents are exactly those of Billing's turnover (« Facturé » of the
forecast): issued invoices (service HT, at their invoice date, D9), issued credit notes not cancelled (negative, « Note de crédit jj/mm », Zurich day of
issuance) and VALIDATED historical invoices (at their invoice date). NB J = days of these documents; TOTAL = sum of the row's document columns;
Statut « TERMINE » for a closed mandate. Rows: the mandates with a document in the year, and in the sheet of the current year every open mandate.
"""
import io
from datetime import date, datetime
from decimal import Decimal

from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

from app import billing_rules as rules

NAMES_2025 = ["JAN", "FEV", "MAR", "AVR", "MAY", "JUN", "JULY", "AUG", "SEPT", "OCT", "NOV", "DEC"]
NAMES_FROM_2026 = ["JAN", "FEB", "MAR", "APR", "MAY", "JUN", "JUL", "AUG", "SEPT", "OCT", "NOV", "DEC"]
HEAD = ["Date(s)", "Motif", "Projet", "Numéro de projet", "Heure début", "Heure fin", "Heures de travail", "Colonne1"]
RECAP = ["PROJET", "TOTAL H", "TOTAL J", "Num projet", "PO", "Jours PO", "J mois précédant(s)", "NB J TOTAL", "J restants sur PO"]
WIDTHS = {"A": 2.7, "B": 32.9, "C": 31.6, "D": 21.3, "E": 16.9, "F": 20.7, "G": 19.1, "H": 18.7, "I": 12.0, "J": 36.6, "K": 11.6, "L": 12.9, "M": 16.9,
          "N": 24.0, "O": 9.1, "P": 20.7, "Q": 18.3, "R": 22.3}
HOURS_FMT, TOTAL_FMT, DATE_FMT, MONTH_FMT = "0.00", "#,##0.00_ ;\\-#,##0.00\\ ", "mm-dd-yy", "mmm-yy"
TABLE_STYLE = TableStyleInfo(name="TableStyleLight18", showRowStripes=True)
SUIVI_HEAD = ["PROJET", "Dates", "Client", "Num projet", "PO", "NB J PO", "NB J ", "TARIF", "TOTAL"]
SUIVI_WIDTHS = {"A": 24.3, "B": 18.0, "C": 11.0, "D": 18.1, "E": 28.0, "F": 9.0, "G": 9.0, "H": 15.7, "I": 15.7}
CHF_FMT = '_-* #,##0.00\\ [$CHF-100C]_-;\\-* #,##0.00\\ [$CHF-100C]_-;_-* "-"??\\ [$CHF-100C]_-;_-@_-'


def sheet_name(y, m):
    return NAMES_2025[m - 1] if y <= 2025 else f"{NAMES_FROM_2026[m - 1]} {y % 100:02d}"


def decimal_hours(clock):
    """« 14:09 » → 14.15 (hours + minutes / 60), the notation of the workbook (A-03)."""
    h, mi = clock.split(":")
    return round(int(h) + int(mi) / 60, 6)


def put(ws, r, col, v):
    """A value typed by a user (motif, mandate, client, PO…) written as TEXT: « =… » never becomes a formula (a broken file, or a live formula in
    Google Sheets), and the control characters a .xlsx cannot hold (pasted from Word) are dropped instead of failing the whole export."""
    if isinstance(v, str): v = ILLEGAL_CHARACTERS_RE.sub("", v)
    cell = ws.cell(r, col, v if v not in ("", None) else None)
    if isinstance(cell.value, str): cell.data_type = "s"
    return cell


def entry_cells(e):
    """(start, end, hours, formula?) of one entry: the formula fin − début only when it gives exactly the hours stored in Billing."""
    if e["start_time"] and e["end_time"]:
        s, f = decimal_hours(e["start_time"]), decimal_hours(e["end_time"])
        return s, f, e["hours"], abs((f - s) - e["hours"]) < 1e-6
    return None, None, e["hours"], False


def month_of(e, placement):
    """Sheet of an entry: its month of service (« month »), or for an imported entry the sheet it came from when that sheet name is known (« source »)."""
    y, m = int(e["work_date"][:4]), int(e["work_date"][5:7])
    if placement == "source" and e["source"] == "historical_excel" and "!" in (e["source_ref"] or ""):
        src = e["source_ref"].split("!")[0]
        for yy in range(y - 1, y + 2):
            for mm in range(1, 13):
                if sheet_name(yy, mm) == src: return yy, mm
    return y, m


def build(conn, placement="month", today=None):
    """The workbook as bytes, and a summary {sheets, entries, hours} to check it against Billing."""
    rows = [dict(r) for r in conn.execute("""SELECT e.*,p.name project,p.project_number,p.po,p.ordered_days,p.day_hours FROM entries e
                                             JOIN projects p ON p.id=e.project_id ORDER BY e.work_date,e.start_time,e.id""")]
    months = {}
    for e in rows: months.setdefault(month_of(e, placement), []).append(e)
    wb = Workbook(); wb.remove(wb.active)
    recap_rows, before_days = {}, {}          # previous sheet's recap rows {project id: (sheet, row)}; days of each mandate before the month
    total_hours = 0.0
    for n, (y, m) in enumerate(sorted(months)):
        es = sorted(months[(y, m)], key=lambda e: (e["work_date"], e["start_time"] or "", e["id"]))
        name = sheet_name(y, m); ws = wb.create_sheet(name); tname = f"FeuilleDeTemps_{y}_{m:02d}"
        for col, w in WIDTHS.items(): ws.column_dimensions[col].width = w
        ws["B1"] = "Feuille de temps"; ws["B1"].font = Font(size=16, bold=True)
        ws["B2"] = datetime(y, m, 1); ws["B2"].number_format = MONTH_FMT
        ws["B3"] = "Total des heures \nde travail du mois"; ws["B3"].alignment = Alignment(wrap_text=True)
        last = 5 + len(es)
        ws["B4"] = f"=SUM(H6:H{last})"; ws["B4"].number_format = TOTAL_FMT
        for i, h in enumerate(HEAD): ws.cell(5, 2 + i, h)
        for r, e in enumerate(es, start=6):
            s, f, hours, formula = entry_cells(e)
            d = date.fromisoformat(e["work_date"])
            ws.cell(r, 2, datetime(d.year, d.month, d.day)).number_format = DATE_FMT
            put(ws, r, 3, e["description"]); put(ws, r, 4, e["project"]); put(ws, r, 5, e["project_number"])
            if s is not None:
                ws.cell(r, 6, s).number_format = HOURS_FMT; ws.cell(r, 7, f).number_format = HOURS_FMT
            ws.cell(r, 8, f"=G{r}-F{r}" if formula else hours).number_format = HOURS_FMT
            total_hours += hours
        tot = last + 1
        ws.cell(tot, 2, "Total"); ws.cell(tot, 8, f"=SUBTOTAL(109,H6:H{last})").number_format = HOURS_FMT
        ws.add_table(Table(displayName=tname, ref=f"B5:I{tot}", totalsRowCount=1, tableStyleInfo=TABLE_STYLE))
        # recap: one row per mandate with hours this month
        per = {}
        for e in es: per.setdefault(e["project_id"], {"e": e, "h": 0.0})["h"] += e["hours"]
        for i, h in enumerate(RECAP): ws.cell(6, 10 + i, h)
        mine = {}
        for r, (pid, x) in enumerate(sorted(per.items(), key=lambda kv: kv[1]["e"]["project"]), start=7):
            e, dh = x["e"], x["e"]["day_hours"] or 8
            put(ws, r, 10, e["project"]); ws.cell(r, 11, round(x["h"], 6)).number_format = HOURS_FMT
            ws.cell(r, 12, f"=K{r}/{dh:g}").number_format = HOURS_FMT
            put(ws, r, 13, e["project_number"]); put(ws, r, 14, e["po"]); ws.cell(r, 15, e["ordered_days"] or None)
            prev = recap_rows.get(pid)
            ws.cell(r, 16, f"='{prev[0]}'!Q{prev[1]}" if prev else round(before_days.get(pid, 0.0), 6)).number_format = HOURS_FMT
            ws.cell(r, 17, f"=P{r}+L{r}").number_format = HOURS_FMT
            ws.cell(r, 18, f"=O{r}-Q{r}").number_format = HOURS_FMT
            mine[pid] = (name, r)
        if per: ws.add_table(Table(displayName=f"Recap_{y}_{m:02d}", ref=f"J6:R{6 + len(per)}", tableStyleInfo=TABLE_STYLE))
        for pid, x in per.items(): before_days[pid] = before_days.get(pid, 0.0) + x["h"] / (x["e"]["day_hours"] or 8)
        recap_rows = mine
    suivi = write_suivi(wb, conn, sorted(months), today or rules.local_today())
    if not wb.sheetnames: wb.create_sheet("Feuille de temps")
    out = io.BytesIO(); wb.save(out)
    return out.getvalue(), {"sheets": wb.sheetnames, "entries": len(rows), "hours": round(total_hours, 6), "suivi": suivi}


def _credited_lines(r):
    """rc38: the manual lines credited by a credit note row, as two lines for rules.amounts (the credit counts in the turnover like the invoice lines did, D6 = A)."""
    return [{"qty": 1, "unit_rate": r[k] or 0, "vat_mode": m} for k, m in (("lines_taxed", "soumis"), ("lines_untaxed", "hors_tva")) if r[k]]


def documents(conn):
    """The documents of Billing's turnover, as the forecast counts them: [{project_id, date, kind ('invoice'|'credit'|'historical'), ht (Decimal), days}]."""
    out = []
    mlines = rules.invoice_lines_by_id(conn)
    for r in conn.execute("""SELECT i.id,i.project_id,i.invoice_date,i.issued_at,i.qty,i.unit,i.unit_rate,p.day_hours FROM invoices i JOIN projects p ON p.id=i.project_id
                             WHERE i.status='issued'"""):
        # rc37 (D6 = A): the manual lines count in the turnover (HT); they are not days of the mandate
        out.append({"project_id": r["project_id"], "date": rules.doc_date(r), "kind": "invoice", "ht": rules.amounts(r["qty"], r["unit_rate"], 0, lines=mlines.get(r["id"], ()))["subtotal"],
                    "days": rules.days(r["qty"], r["unit"], r["day_hours"])})
    for r in conn.execute("""SELECT i.project_id,cn.issued_at,cn.qty,cn.amount,cn.lines_taxed,cn.lines_untaxed,i.unit,p.day_hours FROM credit_notes cn JOIN invoices i ON i.id=cn.invoice_id
                             JOIN projects p ON p.id=i.project_id WHERE cn.status='issued' AND cn.id NOT IN (SELECT credit_note_id FROM credit_note_cancellations)"""):
        out.append({"project_id": r["project_id"], "date": rules.local_date(r["issued_at"]), "kind": "credit", "ht": -rules.amounts(1, r["amount"], 0, lines=_credited_lines(r))["subtotal"],
                    "days": -rules.days(r["qty"], r["unit"], r["day_hours"]) if r["qty"] is not None else 0.0})
    for r in conn.execute("""SELECT project_id,invoice_date,qty_days,amount_ht FROM historical_invoices WHERE status='validated' AND invoice_date IS NOT NULL
                             AND amount_ht IS NOT NULL"""):
        out.append({"project_id": r["project_id"], "date": r["invoice_date"], "kind": "historical", "ht": rules.amounts(1, r["amount_ht"], 0)["ht"],
                    "days": r["qty_days"] or 0.0})
    return out


def column_label(kind, d):
    return f"{'Note de crédit' if kind == 'credit' else 'Facture'} {d[8:10]}/{d[5:7]}"


def write_suivi(wb, conn, months, today):
    """One SUIVI sheet per year (see the module docstring), inserted after the year's last month sheet. Returns {year: turnover HT as text}."""
    projects = {r["id"]: dict(r) for r in conn.execute("SELECT * FROM projects")}
    docs = documents(conn)
    years = {int(d["date"][:4]) for d in docs}
    if any(p["active"] for p in projects.values()): years.add(today.year)
    turnover = {}
    for k, y in enumerate(sorted(years)):
        mine = [d for d in docs if int(d["date"][:4]) == y]
        cols = sorted({(d["date"], d["kind"] == "credit") for d in mine})
        labels = [column_label("credit" if cr else "invoice", dt) for dt, cr in cols]
        rows = {}
        for d in mine: rows.setdefault(d["project_id"], []).append(d)
        if y == today.year:
            for pid, p in projects.items():
                if p["active"]: rows.setdefault(pid, [])
        order = sorted(rows, key=lambda pid: (min((d["date"] for d in rows[pid]), default="9999"), projects[pid]["name"]))
        ws = wb.create_sheet(f"SUIVI {y}", sum(1 for yy, _ in months if yy <= y) + k)
        for col, w in SUIVI_WIDTHS.items(): ws.column_dimensions[col].width = w
        head = SUIVI_HEAD + labels + ["Statut"]
        for i in range(len(SUIVI_HEAD), len(head)): ws.column_dimensions[get_column_letter(i + 1)].width = 17.3
        ws["A1"] = y; ws["A1"].font = Font(bold=True)
        for i, h in enumerate(head): ws.cell(2, i + 1, h).font = Font(bold=True)
        first, last_col = get_column_letter(len(SUIVI_HEAD) + 1), get_column_letter(len(SUIVI_HEAD) + len(cols))
        for r, pid in enumerate(order, start=3):
            p = projects[pid]
            for i, v in enumerate((p["name"], p["dates_label"], p["client"], p["project_number"], p["po"], p["ordered_days"]), start=1): put(ws, r, i, v)
            ws.cell(r, 7, round(sum(d["days"] for d in rows[pid]), 6)).number_format = "0.00"
            ws.cell(r, 8, p["rate"]).number_format = CHF_FMT
            ws.cell(r, 9, f"=SUM({first}{r}:{last_col}{r})" if cols else 0).number_format = CHF_FMT
            for j, (dt, cr) in enumerate(cols, start=len(SUIVI_HEAD) + 1):
                amount = sum((d["ht"] for d in rows[pid] if d["date"] == dt and (d["kind"] == "credit") == cr), Decimal(0))
                if any(d["date"] == dt and (d["kind"] == "credit") == cr for d in rows[pid]): ws.cell(r, j, float(amount)).number_format = CHF_FMT
            if not p["active"]: ws.cell(r, len(head), "TERMINE")
        last = 2 + len(order); tot = last + 1
        if order:
            ws.add_table(Table(displayName=f"Suivi_{y}", ref=f"A2:{get_column_letter(len(head))}{last}", tableStyleInfo=TABLE_STYLE))
            ws.cell(tot, 8, "CA").font = Font(bold=True)
            ws.cell(tot, 9, f"=SUBTOTAL(109,I3:I{last})").number_format = CHF_FMT
            for j in range(len(SUIVI_HEAD) + 1, len(SUIVI_HEAD) + len(cols) + 1):
                L = get_column_letter(j); ws.cell(tot, j, f"=SUM({L}3:{L}{last})").number_format = CHF_FMT
            for j in range(8, len(SUIVI_HEAD) + len(cols) + 1): ws.cell(tot, j).font = Font(bold=True)
        turnover[str(y)] = str(sum((d["ht"] for d in mine), Decimal(0)))
    return turnover
